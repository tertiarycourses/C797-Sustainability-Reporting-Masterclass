import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
W, H = p.slide_width, p.slide_height
idx = [int(x) for x in sys.argv[2].split(",")] if len(sys.argv) > 2 else range(len(p.slides))
for i in idx:
    s = p.slides[i]
    print(f"===== slide {i} layout={s.slide_layout.name}")
    for sh in s.shapes:
        pos = f"{sh.left/W:.2f},{sh.top/H:.2f}" if sh.left is not None else "?"
        if sh.has_text_frame and sh.text_frame.text.strip():
            paras = [ "|".join(r.text for r in pa.runs) for pa in sh.text_frame.paragraphs]
            print(f"  [{sh.shape_id}] {sh.shape_type} @{pos} :: " + " // ".join(paras))
        elif getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                print(f"  [{sh.shape_id}] TABLE :: " + " | ".join(c.text.replace("\n"," / ") for c in r.cells))
        elif sh.shape_type == 13:
            print(f"  [{sh.shape_id}] PICTURE @{pos} {sh.name}")
        elif sh.shape_type == 6:
            print(f"  [{sh.shape_id}] GROUP @{pos}")
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        print("  NOTES:", s.notes_slide.notes_text_frame.text[:300].replace("\n"," / "))
