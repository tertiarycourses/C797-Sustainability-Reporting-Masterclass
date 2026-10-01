"""Post-conversion fix-ups for the C797 deck: strip the WSQ layer convert_deck.py left behind."""
import re
from pptx import Presentation

PATH = "courseware/Sustainability Reporting Masterclass (C797)-v1.0.pptx"
p = Presentation(PATH)


def set_para(para, line):
    runs = [r for r in para.runs if r.text] or para.runs
    runs[0].text = line
    for r in para.runs:
        if r._r is not runs[0]._r:
            r.text = ""


def paras(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                yield sh, para, "".join(r.text for r in para.runs)


KA = re.compile(r"\s*•\s*[KA]\d+(?:,\s*[KA]\d+)*(?=\s*•)")   # "LU1 • K1, K2, A1 • MODEL" -> "LU1 • MODEL"
SUBS = {
    "Trainer / assessor / course administrator": "Trainer / course administrator",
    "Retain the approved attendance or assessment record.": "Retain the approved attendance record.",
    "Day 1: LU1–2 • Day 2: LU3–4 • WA + CS": "Day 1: LU1–2 • Day 2: LU3–4",
    "Submit TRAQOM and confirm receipt.": "Complete the course feedback survey before you leave.",
}
SRCNOTE = re.compile(r"^(Original slide \d+) from “WSQ - Master Trainer Slides - .*”\.$")

# ---- cover code/version line: convert_deck.py flattened its <a:br/> into a literal "_x000B_";
# restore the source text body and edit only the run texts.
import copy
SRC = "source-wsq/courseware/Sustainability-Reporting-and-Engineering-Strategies-for-Senior-Leaders-v5.0.pptx"
src_sh = next(sh for sh in Presentation(SRC).slides[0].shapes if sh.shape_id == 4)
dst_sh = next(sh for sh in p.slides[0].shapes if sh.shape_id == 4)
dst_sh._element.replace(dst_sh._element.txBody, copy.deepcopy(src_sh._element.txBody))
for r in dst_sh.text_frame.paragraphs[0].runs:
    r.text = r.text.replace("TGS-2025052343", "C797").replace("Version 5.0", "Version 1.0")
assert "_x000B_" not in dst_sh.text_frame.text and "TGS" not in dst_sh.text_frame.text

hits = 0
for s in p.slides:
    for sh, para, t in paras(s):
        new = SUBS.get(t, t)
        new = KA.sub("", new)
        new = SRCNOTE.sub(r"\1 from the master trainer slides.", new)
        if new != t:
            set_para(para, new)
            hits += 1

p.core_properties.title = "Sustainability Reporting Masterclass"
p.core_properties.subject = "C797"
p.save(PATH)
print("saved", hits, "paragraphs rewritten")
